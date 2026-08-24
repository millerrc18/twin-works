import openpyxl, json, datetime
src = r"C:/Users/ryan.c.miller/Downloads/Copy of GAC RTG 2026.xlsx"
wb = openpyxl.load_workbook(src, data_only=True)
ws = wb["Sheet1"]

def d(v):
    if isinstance(v,(datetime.datetime,datetime.date)):
        return v.strftime("%Y-%m-%d")
    return None

# Column map (1-based): B=2 date, C=3 unit, D=4 target, E=5 AJ/LAM, F=6 A1/ASSY, G=7 A2/PAINT, H=8 FS/SHIP, I=9 SL/PO, J=10 SL(radome)
def grab(r):
    return dict(
        unit=ws.cell(r,3).value,
        target=d(ws.cell(r,4).value),
        m1=d(ws.cell(r,5).value),
        m2=d(ws.cell(r,6).value),
        m3=d(ws.cell(r,7).value),
        m4=d(ws.cell(r,8).value),
        colI=ws.cell(r,9).value,
        colJ=ws.cell(r,10).value,
    )

# LH: rows 5-24 ; RH: rows 30-51 ; Radome: rows 54-92 (373=r54 ... 547=r92); skip 548/549
lh  = [grab(r) for r in range(5,25)]
rh  = [grab(r) for r in range(30,52)]
rad = [grab(r) for r in range(54,93)]

out = {"lh":lh,"rh":rh,"rad":rad}
with open(r"C:/Users/ryan.c.miller/Downloads/rtg-tracker-build/rtg_data.json","w") as f:
    json.dump(out,f,indent=1)
print("LH",len(lh),"RH",len(rh),"RAD",len(rad))
print("LH sample:",lh[0], lh[4])
print("RAD sample:",rad[0], rad[-1])
